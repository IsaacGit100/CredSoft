"""
Shared utilities for all reports.
"""

from datetime import datetime
from decimal import Decimal
from io import BytesIO
from urllib.parse import urlencode

from django.db.models import Q, Sum
from django.http import HttpResponse
from django.template.loader import render_to_string

from django_ledger.models import JournalEntryModel, TransactionModel

ROOT_ROLES = {
    "root_coa",
    "root_assets",
    "root_liabilities",
    "root_capital",
    "root_income",
    "root_cogs",
    "root_expenses",
}
CASH_CODES = ["1010", "1020"]

def parse_date(value):
    """Accept dd/mm/yyyy, dd-mm-yyyy, yyyy-mm-dd. Returns date or None."""
    if not value:
        return None
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(value, fmt).date()
        except ValueError:
            continue
    return None


def period_label(start, end):
    if start and end:
        return f"{start:%d/%m/%Y} to {end:%d/%m/%Y}"
    if start:
        return f"From {start:%d/%m/%Y}"
    if end:
        return f"Up to {end:%d/%m/%Y}"
    return "All dates"


def _base_qs(entity, start=None, end=None):
    qs = TransactionModel.objects.filter(
        journal_entry__ledger__entity=entity,
        journal_entry__posted=True,
    )
    if start:
        qs = qs.filter(journal_entry__timestamp__date__gte=start)
    if end:
        qs = qs.filter(journal_entry__timestamp__date__lte=end)
    return qs


def get_trial_balance(entity, start=None, end=None):
    """
    Return (rows, total_dr, total_cr).
    Each row: code, name, balance_type, role, dr, cr, balance.
    """
    qs = _base_qs(entity, start, end)

    raw = (
        qs.values(
            "account__code",
            "account__name",
            "account__balance_type",
            "account__role",
        )
        .annotate(
            dr=Sum("amount", filter=Q(tx_type="debit")),
            cr=Sum("amount", filter=Q(tx_type="credit")),
        )
        .order_by("account__code")
    )

    rows = []
    total_dr = Decimal("0")
    total_cr = Decimal("0")

    for r in raw:
        if r["account__role"] in ROOT_ROLES:
            continue

        dr = r["dr"] or Decimal("0")
        cr = r["cr"] or Decimal("0")
        if dr == 0 and cr == 0:
            continue

        balance_type = r["account__balance_type"] or "debit"
        balance = (dr - cr) if balance_type == "debit" else (cr - dr)

        rows.append(
            {
                "code": r["account__code"],
                "name": r["account__name"],
                "balance_type": balance_type,
                "role": r["account__role"],
                "dr": dr,
                "cr": cr,
                "balance": balance,
            }
        )
        total_dr += dr
        total_cr += cr

    return rows, total_dr, total_cr


def clean_querystring(request):
    """Query string with format/download stripped — for the View/PDF/Excel links."""
    qs = request.GET.copy()
    qs.pop("format", None)
    qs.pop("download", None)
    return qs.urlencode()


# ---------- Exports ----------


def render_pdf(request, template_name, context, filename="report.pdf"):
    from xhtml2pdf import pisa

    html = render_to_string(template_name, context, request=request)
    buffer = BytesIO()
    result = pisa.CreatePDF(src=html, dest=buffer, encoding="utf-8")

    if result.err:
        return HttpResponse("PDF generation failed", status=500)

    response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
    mode = "attachment" if request.GET.get("download") == "1" else "inline"
    response["Content-Disposition"] = f'{mode}; filename="{filename}"'
    return response


def render_excel(
    headers,
    rows,
    filename="report.xlsx",
    sheet_name="Report",
    title=None,
    subtitle=None,
):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]

    row_idx = 1

    if title:
        c = ws.cell(row=row_idx, column=1, value=title)
        c.font = Font(bold=True, size=14)
        c.alignment = Alignment(horizontal="center")
        ws.merge_cells(
            start_row=row_idx, start_column=1, end_row=row_idx, end_column=len(headers)
        )
        row_idx += 1

    if subtitle:
        c = ws.cell(row=row_idx, column=1, value=subtitle)
        c.font = Font(italic=True, size=11, color="666666")
        c.alignment = Alignment(horizontal="center")
        ws.merge_cells(
            start_row=row_idx, start_column=1, end_row=row_idx, end_column=len(headers)
        )
        row_idx += 1

    row_idx += 1

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1A56DB")
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=row_idx, column=col, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
    row_idx += 1

    for row in rows:
        for col, val in enumerate(row, start=1):
            ws.cell(row=row_idx, column=col, value=val)
        row_idx += 1

    for col in range(1, len(headers) + 1):
        letter = get_column_letter(col)
        maxlen = max(
            [len(str(headers[col - 1]))]
            + [len(str(r[col - 1])) if col - 1 < len(r) else 0 for r in rows]
        )
        ws.column_dimensions[letter].width = min(max(maxlen + 2, 12), 45)

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


def as_at_label(end, fallback_today=True):
    """For balance-sheet-style reports: 'as at 30/09/2026'."""
    from datetime import date

    if end:
        return f"as at {end:%d/%m/%Y}"
    if fallback_today:
        return f"as at {date.today():%d/%m/%Y}"
    return ""


def for_period_label(start, end):
    """For flow reports: 'for the period 01/09/2026 to 30/09/2026'."""
    if start and end:
        return f"for the period {start:%d/%m/%Y} to {end:%d/%m/%Y}"
    if start:
        return f"from {start:%d/%m/%Y}"
    if end:
        return f"up to {end:%d/%m/%Y}"
    return "for all dates"


## ======================================Journals=========================
def get_journal_list(entity, start=None, end=None):
    """
    Return (lines, total_dr, total_cr).

    One entry per journal line (each JE has 2 lines: DR + CR).
    Ordered by date, then JE number, then debit-before-credit.
    """
    qs = (
        _base_qs(entity, start, end)
        .select_related("account", "journal_entry")
        .order_by(
            "journal_entry__timestamp",
            "journal_entry__je_number",
            "-tx_type",  # 'debit' sorts before 'credit'
            "account__code",
        )
    )

    lines = []
    total_dr = Decimal("0")
    total_cr = Decimal("0")

    for tx in qs:
        je = tx.journal_entry
        is_dr = tx.tx_type == "debit"
        amount = tx.amount or Decimal("0")

        if is_dr:
            total_dr += amount
        else:
            total_cr += amount

        lines.append(
            {
                "date": je.timestamp.date(),
                "je_number": je.je_number,
                "je_uuid": str(je.uuid),
                "je_description": je.description or "",
                "account_code": tx.account.code,
                "account_name": tx.account.name,
                "description": tx.description or je.description or "",
                "dr": amount if is_dr else None,
                "cr": amount if not is_dr else None,
                "is_debit": is_dr,
            }
        )

    return lines, total_dr, total_cr


def get_income_statement(entity, start=None, end=None):
    """
    Return dict with:
        revenue_rows, total_revenue
        cogs_rows,    total_cogs
        expense_rows, total_expense
        gross_profit  (revenue - cogs)
        net_income    (gross_profit - expense)
    """
    qs = _base_qs(entity, start, end)

    raw = (
        qs.values(
            "account__code",
            "account__name",
            "account__role",
            "account__balance_type",
        )
        .annotate(
            dr=Sum("amount", filter=Q(tx_type="debit")),
            cr=Sum("amount", filter=Q(tx_type="credit")),
        )
        .order_by("account__code")
    )

    revenue_rows, cogs_rows, expense_rows = [], [], []
    total_revenue = Decimal("0")
    total_cogs = Decimal("0")
    total_expense = Decimal("0")

    for r in raw:
        role = r["account__role"]
        if role not in ("revenue", "expense", "cogs"):
            continue

        dr = r["dr"] or Decimal("0")
        cr = r["cr"] or Decimal("0")
        if dr == 0 and cr == 0:
            continue

        # Revenue: credit-normal → amount = cr - dr
        # Expense/COGS: debit-normal → amount = dr - cr
        if role == "revenue":
            amount = cr - dr
        else:
            amount = dr - cr

        entry = {
            "code": r["account__code"],
            "name": r["account__name"],
            "amount": amount,
        }

        if role == "revenue":
            revenue_rows.append(entry)
            total_revenue += amount
        elif role == "cogs":
            cogs_rows.append(entry)
            total_cogs += amount
        else:  # expense
            expense_rows.append(entry)
            total_expense += amount

    gross_profit = total_revenue - total_cogs
    net_income = gross_profit - total_expense

    return {
        "revenue_rows": revenue_rows,
        "cogs_rows": cogs_rows,
        "expense_rows": expense_rows,
        "total_revenue": total_revenue,
        "total_cogs": total_cogs,
        "total_expense": total_expense,
        "gross_profit": gross_profit,
        "net_income": net_income,
    }


def get_balance_sheet(entity, end=None):
    """
    Balance Sheet as at a date. Cumulative — ignores `start`.

    Returns dict:
        asset_rows, total_assets
        liability_rows, total_liabilities
        equity_rows, total_equity_accounts
        current_earnings       (net income for fiscal-year-to-date;
                                placeholder = all-time net income for now)
        total_equity           = equity accounts + current earnings
        total_liabilities_equity
        balanced               (total_assets == total_liabilities_equity)
    """
    # Assets, Liabilities, Equity — cumulative to `end`
    qs = _base_qs(entity, None, end)

    raw = (
        qs.values(
            "account__code",
            "account__name",
            "account__role",
            "account__balance_type",
        )
        .annotate(
            dr=Sum("amount", filter=Q(tx_type="debit")),
            cr=Sum("amount", filter=Q(tx_type="credit")),
        )
        .order_by("account__code")
    )

    asset_rows, liability_rows, equity_rows = [], [], []
    total_assets = Decimal("0")
    total_liabilities = Decimal("0")
    total_equity_accounts = Decimal("0")

    for r in raw:
        role = r["account__role"]
        if role not in ("asset", "liability", "equity"):
            continue

        dr = r["dr"] or Decimal("0")
        cr = r["cr"] or Decimal("0")
        if dr == 0 and cr == 0:
            continue

        balance_type = r["account__balance_type"] or "debit"
        # Assets: dr - cr  (contra-assets like accumulated depreciation
        # naturally come out negative, which is correct)
        # Liabilities/Equity: cr - dr
        if role == "asset":
            amount = dr - cr
        else:
            amount = cr - dr

        entry = {
            "code": r["account__code"],
            "name": r["account__name"],
            "amount": amount,
        }

        if role == "asset":
            asset_rows.append(entry)
            total_assets += amount
        elif role == "liability":
            liability_rows.append(entry)
            total_liabilities += amount
        else:
            equity_rows.append(entry)
            total_equity_accounts += amount

    # Current-period earnings — from income statement (all time for now)
    is_data = get_income_statement(entity, None, end)
    current_earnings = is_data["net_income"]

    total_equity = total_equity_accounts + current_earnings
    total_liabilities_equity = total_liabilities + total_equity
    balanced = total_assets == total_liabilities_equity

    return {
        "asset_rows": asset_rows,
        "liability_rows": liability_rows,
        "equity_rows": equity_rows,
        "total_assets": total_assets,
        "total_liabilities": total_liabilities,
        "total_equity_accounts": total_equity_accounts,
        "current_earnings": current_earnings,
        "total_equity": total_equity,
        "total_liabilities_equity": total_liabilities_equity,
        "balanced": balanced,
    }


def _cash_balance(entity, before=None, up_to=None):
    """Total cash balance (DR - CR) at a point in time."""
    qs = TransactionModel.objects.filter(
        journal_entry__ledger__entity=entity,
        journal_entry__posted=True,
        account__code__in=CASH_CODES,
    )
    if before:
        qs = qs.filter(journal_entry__timestamp__date__lt=before)
    if up_to:
        qs = qs.filter(journal_entry__timestamp__date__lte=up_to)

    dr = qs.filter(tx_type="debit").aggregate(s=Sum("amount"))["s"] or Decimal("0")
    cr = qs.filter(tx_type="credit").aggregate(s=Sum("amount"))["s"] or Decimal("0")
    return dr - cr


def _classify_cash_entry(non_cash_line):
    """
    Given the non-cash side of a JE, decide the section.
    Returns 'operating' | 'investing' | 'financing' | None
    """
    role = non_cash_line.account.role
    code = non_cash_line.account.code

    if role in ("revenue", "expense", "cogs"):
        return "operating"
    if role in ("liability", "equity"):
        return "financing"
    if role == "asset":
        # PPE / non-current → investing; current assets (AR, Inventory, Prepaid) → operating
        if code.startswith("11") or code.startswith("109"):
            return "investing"
        return "operating"
    return None


def get_cash_flow(entity, start=None, end=None):
    """
    Direct-method Cash Flow based on cash-account movements in the period.

    Returns dict:
        operating, investing, financing   (lists of line dicts)
        total_operating, total_investing, total_financing
        net_change
        opening_cash, closing_cash
        reconciled   (opening + net_change == closing)
    """
    from collections import defaultdict

    qs = _base_qs(entity, start, end).select_related("account", "journal_entry")

    by_je = defaultdict(list)
    for tx in qs:
        by_je[tx.journal_entry_id].append(tx)

    operating, investing, financing = [], [], []

    for je_id, lines in by_je.items():
        cash_lines = [l for l in lines if l.account.code in CASH_CODES]
        non_cash_lines = [l for l in lines if l.account.code not in CASH_CODES]

        if not cash_lines:
            continue

        cash_effect = Decimal("0")
        for l in cash_lines:
            if l.tx_type == "debit":
                cash_effect += l.amount or Decimal("0")
            else:
                cash_effect -= l.amount or Decimal("0")

        if not non_cash_lines:
            continue

        ref = non_cash_lines[0]
        section = _classify_cash_entry(ref)
        if section is None:
            continue

        entry = {
            "date": lines[0].journal_entry.timestamp.date(),
            "je_number": lines[0].journal_entry.je_number,
            "account_code": ref.account.code,
            "account_name": ref.account.name,
            "description": (
                ref.description or lines[0].journal_entry.description or ""
            )[:100],
            "amount": cash_effect,
        }

        if section == "operating":
            operating.append(entry)
        elif section == "investing":
            investing.append(entry)
        else:
            financing.append(entry)

    # Sort each section by date
    operating.sort(key=lambda e: (e["date"], e["je_number"]))
    investing.sort(key=lambda e: (e["date"], e["je_number"]))
    financing.sort(key=lambda e: (e["date"], e["je_number"]))

    total_operating = sum((e["amount"] for e in operating), Decimal("0"))
    total_investing = sum((e["amount"] for e in investing), Decimal("0"))
    total_financing = sum((e["amount"] for e in financing), Decimal("0"))
    net_change = total_operating + total_investing + total_financing

    opening_cash = _cash_balance(entity, before=start) if start else Decimal("0")
    closing_cash = _cash_balance(entity, up_to=end)

    return {
        "operating": operating,
        "investing": investing,
        "financing": financing,
        "total_operating": total_operating,
        "total_investing": total_investing,
        "total_financing": total_financing,
        "net_change": net_change,
        "opening_cash": opening_cash,
        "closing_cash": closing_cash,
        "reconciled": (opening_cash + net_change) == closing_cash,
    }


def get_trans_records(entity, start=None, end=None, trans_type=None, sub_module=None):
    """
    Raw Trans rows filtered by date. Optional trans_type and sub_module.
    """
    from RecPayApp.models import Trans

    qs = Trans.objects.filter(entity=entity).order_by("date", "id")
    if start:
        qs = qs.filter(date__gte=start)
    if end:
        qs = qs.filter(date__lte=end)
    if trans_type:
        qs = qs.filter(trans_type=trans_type)
    if sub_module:
        qs = qs.filter(sub_module=sub_module)

    lines = []
    total_receipts = Decimal("0")
    total_payments = Decimal("0")

    for t in qs:
        amount = t.amount or Decimal("0")
        if t.trans_type == "Receipts":
            total_receipts += amount
        elif t.trans_type == "Payments":
            total_payments += amount

        # Party name — whichever is set
        if t.church_member:
            party = t.church_member.full_name
        elif t.member:
            party = t.member.full_name
        elif t.non_member_name:
            party = t.non_member_name
        else:
            party = ""

        lines.append(
            {
                "id": t.pk,
                "date": t.date,
                "trans_no": t.trans_no or "",
                "rec_vou_no": t.rec_vou_no or "",
                "trans_type": t.trans_type or "",
                "module": t.module or "",
                "sub_module": t.sub_module or "",
                "party": party,
                "ledger_code": t.ledger_code or "",
                "ledger_name": t.ledger_name or "",
                "purpose": t.purpose or "",
                "details": t.details or "",
                "pay_mode": t.pay_mode or "",
                "amount": amount,
                "journal_status": t.journal_status or "",
                "journal_entry_id": t.journal_entry_id or "",
            }
        )

    return {
        "lines": lines,
        "total_receipts": total_receipts,
        "total_payments": total_payments,
        "net": total_receipts - total_payments,
    }


def get_journal_records(entity, start=None, end=None, account_query=None):
    """
    Journal lines filtered by date and by account (code or name, icontains).
    """
    qs = (
        _base_qs(entity, start, end)
        .select_related("account", "journal_entry")
        .order_by(
            "journal_entry__timestamp",
            "journal_entry__je_number",
            "account__code",
        )
    )

    if account_query:
        qs = qs.filter(
            Q(account__code__icontains=account_query)
            | Q(account__name__icontains=account_query)
        )

    lines = []
    total_dr = Decimal("0")
    total_cr = Decimal("0")

    for tx in qs:
        je = tx.journal_entry
        is_dr = tx.tx_type == "debit"
        amount = tx.amount or Decimal("0")

        if is_dr:
            total_dr += amount
        else:
            total_cr += amount

        lines.append(
            {
                "date": je.timestamp.date(),
                "je_number": je.je_number,
                "je_uuid": str(je.uuid),
                "description": tx.description or je.description or "",
                "account_code": tx.account.code,
                "account_name": tx.account.name,
                "account_role": tx.account.role or "",
                "account_balance_type": tx.account.balance_type or "",
                "tx_type": tx.tx_type,
                "is_debit": is_dr,
                "dr": amount if is_dr else None,
                "cr": amount if not is_dr else None,
            }
        )

    return {
        "lines": lines,
        "total_dr": total_dr,
        "total_cr": total_cr,
    }


def get_dashboard_data(entity, start=None, end=None):
    """
    Aggregated finance position for the dashboard.

    Returns dict with:
        cash, bank, total_cash
        total_receipts, total_payments, net_flow      (Trans, filtered period)
        total_revenue, total_expense, net_income      (JEs, filtered period)
        total_assets, total_liabilities, total_equity (JEs, cumulative)
        pending_trans, posted_trans                   (Trans counts)
        recent_trans                                  (last 10 Trans rows)
        recent_journals                               (last 5 JEs)
    """
    from RecPayApp.models import Trans

    # ---- Cash position (as at `end` or today) ----
    from datetime import date as _date

    today = _date.today()
    as_at = end or today

    def _balance_for(code):
        qs = TransactionModel.objects.filter(
            journal_entry__ledger__entity=entity,
            journal_entry__posted=True,
            account__code=code,
            journal_entry__timestamp__date__lte=as_at,
        )
        dr = qs.filter(tx_type="debit").aggregate(s=Sum("amount"))["s"] or Decimal("0")
        cr = qs.filter(tx_type="credit").aggregate(s=Sum("amount"))["s"] or Decimal("0")
        return dr - cr

    cash = _balance_for("1010")
    bank = _balance_for("1020")
    total_cash = cash + bank

    # ---- Period flows (Trans) ----
    trans_qs = Trans.objects.filter(entity=entity)
    if start:
        trans_qs = trans_qs.filter(date__gte=start)
    if end:
        trans_qs = trans_qs.filter(date__lte=end)

    total_receipts = trans_qs.filter(trans_type="Receipts").aggregate(s=Sum("amount"))[
        "s"
    ] or Decimal("0")
    total_payments = trans_qs.filter(trans_type="Payments").aggregate(s=Sum("amount"))[
        "s"
    ] or Decimal("0")
    net_flow = total_receipts - total_payments

    # ---- Period Income Statement ----
    is_data = get_income_statement(entity, start, end)

    # ---- Cumulative Balance Sheet ----
    bs_data = get_balance_sheet(entity, end)

    # ---- Trans status ----
    base_trans = Trans.objects.filter(entity=entity)
    pending_trans = base_trans.filter(journal_status="PENDING").count()
    posted_trans = base_trans.filter(journal_status="POSTED").count()

    # ---- Recent activity ----
    recent_trans = list(base_trans.order_by("-date", "-id")[:10])
    recent_journals = JournalEntryModel.objects.filter(
        journal_entry__ledger__entity=entity
    ).order_by("-timestamp", "-created")[:5]

    return {
        "cash": cash,
        "bank": bank,
        "total_cash": total_cash,
        "total_receipts": total_receipts,
        "total_payments": total_payments,
        "net_flow": net_flow,
        "total_revenue": is_data["total_revenue"],
        "total_expense": is_data["total_expense"],
        "net_income": is_data["net_income"],
        "total_assets": bs_data["total_assets"],
        "total_liabilities": bs_data["total_liabilities"],
        "total_equity": bs_data["total_equity"],
        "pending_trans": pending_trans,
        "posted_trans": posted_trans,
        "recent_trans": recent_trans,
        "recent_journals": recent_journals,
    }
